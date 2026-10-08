# Copyright (C) 2026 Sergi Albuixech
# SPDX-License-Identifier: AGPL-3.0-or-later

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT


def _assignment_slots(a):
    """Todas las franjas globales que ocupa una asignación."""
    return a.get("slots", [a["slot"]])


def _exam_range(a):
    """Rango horario real del examen ('09:00-13:30')."""
    label = a.get("slot_label", "")
    return label.split(" ", 1)[1] if " " in label else label


def _duration_label(exam):
    """Duración declarada del examen ('2 h')."""
    try:
        dur = float(exam.get("duration_hours", 2.0) or 2.0)
    except (TypeError, ValueError):
        dur = 2.0
    return f"{dur:g} h"


def export_calendar_to_word(assignment, exams, classrooms, filepath, project_name=""):
    """
    Exporta el calendario generado a un documento Word (.docx).
    
    assignment: lista de asignaciones (cada una con slot, slots, slot_label, exam, classroom, ...)
    """
    doc = Document()
    
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Calibri"
    font.size = Pt(11)

    # ── Portada / Título ──
    title = doc.add_heading(f"📅 Calendario de Exámenes", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if project_name:
        sub = doc.add_paragraph(project_name)
        sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sub.runs[0].font.size = Pt(14)
        sub.runs[0].font.color.rgb = RGBColor(99, 102, 241)
    
    doc.add_paragraph()
    
    # ── Estadísticas ──
    n_exams = len(exams)
    total_students = sum(e["students"] for e in exams)
    n_slots = len({t for a in assignment for t in _assignment_slots(a)})
    n_classrooms = len(classrooms)
    n_studies = len({e["study"] for e in exams})
    n_pcs = sum(int(e.get("computers", 0) or 0) for e in exams)
    pcs_txt = f" · 💻 {n_pcs} ordenadores" if n_pcs else ""
    
    stats = doc.add_paragraph(
        f"📋 {n_exams} exámenes · 🕐 {n_slots} franjas · "
        f"🏫 {n_classrooms} aulas · 👥 {total_students} alumnos · "
        f"📚 {n_studies} estudios{pcs_txt}"
    )
    stats.alignment = WD_ALIGN_PARAGRAPH.CENTER
    stats.runs[0].font.italic = True
    stats.runs[0].font.color.rgb = RGBColor(100, 116, 139)
    
    doc.add_paragraph()
    
    # ── Tabla de asignaciones ──
    doc.add_heading("Distribución de exámenes", level=2)
    
    # Agrupar por slot de inicio
    groups = {}
    for a in sorted(assignment, key=lambda x: (x["slot"], x["classroom"]["name"])):
        groups.setdefault(a["slot"], []).append(a)
    
    for t in sorted(groups.keys()):
        label = groups[t][0]["slot_label"] if groups[t] else f"Franja {t+1}"
        doc.add_heading(f"🕐 {label}", level=3)
        
        table = doc.add_table(rows=1, cols=6)
        table.style = "Light Shading Accent 1"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        hdr = table.rows[0].cells
        hdr[0].text = "Aula"
        hdr[1].text = "Examen"
        hdr[2].text = "Alumnos"
        hdr[3].text = "Estudio"
        hdr[4].text = "Horario"
        hdr[5].text = "Duración"
        
        for x in sorted(groups[t], key=lambda x: x["classroom"]["name"]):
            row = table.add_row().cells
            row[0].text = x["classroom"]["name"]
            row[1].text = x["exam"]["name"]
            row[2].text = str(x["exam"]["students"])
            row[3].text = x["exam"]["study"]
            row[4].text = _exam_range(x)
            row[5].text = _duration_label(x["exam"])
        
        doc.add_paragraph()
    
    # ── Leyenda de estudios ──
    studies = sorted({e["study"] for e in exams})
    doc.add_heading("Estudios", level=2)
    for s in studies:
        count = sum(1 for e in exams if e["study"] == s)
        students = sum(e["students"] for e in exams if e["study"] == s)
        doc.add_paragraph(f"  • {s} — {count} exámenes, {students} alumnos")
    
    # ── Footer ──
    doc.add_paragraph()
    footer = doc.add_paragraph(f"Generado con Generador de Calendario de Exámenes")
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.runs[0].font.size = Pt(9)
    footer.runs[0].font.color.rgb = RGBColor(148, 163, 184)
    
    doc.save(filepath)


def export_exams_to_word(exams, filepath):
    """Exporta la lista de exámenes a un documento Word."""
    doc = Document()
    
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Calibri"
    font.size = Pt(11)
    
    title = doc.add_heading("📋 Lista de Exámenes", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph(f"Total: {len(exams)} exámenes").alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph()
    
    table = doc.add_table(rows=1, cols=8)
    table.style = "Light Shading Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    hdr = table.rows[0].cells
    hdr[0].text = "Nombre"
    hdr[1].text = "Alumnos"
    hdr[2].text = "Estudio"
    hdr[3].text = "Profesor"
    hdr[4].text = "Turno"
    hdr[5].text = "Duración"
    hdr[6].text = "Ordenadores"
    hdr[7].text = "Color"
    
    for e in exams:
        row = table.add_row().cells
        row[0].text = e["name"]
        row[1].text = str(e["students"])
        row[2].text = e["study"]
        row[3].text = e.get("teacher", "")
        shift_map = {"morning": "Mañana", "afternoon": "Tarde"}
        row[4].text = shift_map.get(e.get("preferred_shift", ""), "")
        row[5].text = _duration_label(e)
        pcs = int(e.get("computers", 0) or 0)
        row[6].text = str(pcs) if pcs else "—"
        row[7].text = e.get("color", "")
    
    doc.add_paragraph()
    
    # Resumen por estudio
    doc.add_heading("Resumen por estudio", level=2)
    studies = {}
    for e in exams:
        studies.setdefault(e["study"], []).append(e)
    for s, lst in sorted(studies.items()):
        total = sum(e["students"] for e in lst)
        doc.add_paragraph(f"  • {s} — {len(lst)} exámenes, {total} alumnos")
    
    doc.save(filepath)
